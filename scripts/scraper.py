from pathlib import Path
import re
from collections import defaultdict

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import DataBarRule
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList

# Base URL
BASE_URL = "https://realpython.com/learning-paths/"
# Get project root (parent of scripts directory)
PROJECT_ROOT = Path(__file__).parent.parent
HTML_DIR = PROJECT_ROOT / "html_cache"
OUTPUT_DIR = PROJECT_ROOT / "output"

def get_local_html_path(url):
    """Get the local file path for a URL if it exists"""
    if not HTML_DIR.exists():
        return None
    
    url_clean = url.rstrip('/')
    
    # Handle main learning paths page
    base_clean = BASE_URL.rstrip('/')
    if url_clean == base_clean or url_clean.endswith('/learning-paths'):
        main_file = HTML_DIR / "main_learning_paths.html"
        return main_file if main_file.exists() else None
    
    # Handle learning path pages
    if '/learning-paths/' in url and url.count('/learning-paths/') == 1:
        # Extract learning path name from URL
        # URL format: https://realpython.com/learning-paths/python-basics/
        path_part = url_clean.replace('https://realpython.com/learning-paths/', '').rstrip('/')
        if path_part:
            # Try to find file by matching the path part
            learning_paths_dir = HTML_DIR / "learning_paths"
            if learning_paths_dir.exists():
                # Look for file that matches (could be sanitized filename)
                path_name = path_part.replace('/', '_')
                path_file = learning_paths_dir / f"{path_name}.html"
                if path_file.exists():
                    return path_file
                # Also try to find by searching (in case filename was sanitized differently)
                for file in learning_paths_dir.glob("*.html"):
                    if path_part in file.stem or file.stem in path_part:
                        return file
    
    # Handle resource pages (courses, exercises, quizzes)
    if '/courses/' in url or '/quizzes/' in url or '/tutorials/' in url:
        url_path = url_clean.replace('https://realpython.com/', '').rstrip('/')
        url_path = url_path.replace('/', '_')
        # Handle very long paths
        if len(url_path) > 200:
            url_path = url_path[:200]
        resource_file = HTML_DIR / "resources" / f"{url_path}.html"
        return resource_file if resource_file.exists() else None
    
    return None

def fetch_page(url, use_cache=True):
    """Fetch a page and return BeautifulSoup object. Uses local cache if available."""
    # Try to use local cache first
    if use_cache:
        local_file = get_local_html_path(url)
        if local_file and local_file.exists():
            try:
                soup = BeautifulSoup(local_file.read_text(encoding='utf-8'), 'html.parser')
                return soup
            except Exception as e:
                print(f"⚠️ Error reading cached file {local_file}: {e}, fetching from web...")
    
    # Fall back to fetching from web
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'
    }
    try:
        response = requests.get(url, headers=headers)
        response.raise_for_status()
        return BeautifulSoup(response.content, 'html.parser')
    except requests.RequestException as e:
        print(f"Error fetching {url}: {e}")
        return None

def fetch_main_page():
    """Fetch main learning paths page"""
    soup = fetch_page(BASE_URL)
    if not soup:
        print("⚠️ Could not fetch main learning paths page")
    return soup

def extract_learning_paths(soup):
    """Extract learning path titles and URLs from the main page"""
    if not soup:
        return []

    paths = []
    for card in soup.find_all("learning-path-card"):
        title_tag = card.find("h2")
        link_tag = card.find("a", class_="stretched-link")
        title = title_tag.get_text(strip=True) if title_tag else None
        href = link_tag.get("href") if link_tag else None
        if title and href:
            # Make sure URL is absolute
            if href.startswith('/'):
                href = f"https://realpython.com{href}"
            paths.append({"title": title, "url": href})
    return paths

# Removed explore_single_learning_path - not needed for production

# Removed extract_one_resource - not needed, we use extract_all_resources_from_path instead

def extract_duration_from_resource_page(resource_url, verbose=False):
    """Visit a resource page and extract its duration (top-level and sections)"""
    soup = fetch_page(resource_url)
    if not soup:
        if verbose:
            print("⚠️ Could not fetch resource page")
        return None
    
    # First, try to find top-level duration in the header/hero section
    # It's usually near the title, with "Course duration" title attribute or near "X Lessons"
    top_duration = None
    top_lesson_count = None
    top_level_text = None
    
    # Look for the header area (near h1 title)
    h1 = soup.find('h1')
    if h1:
        # Look in the parent container
        header_area = h1.parent
        if header_area:
            # Look for span with title="Course duration" - this contains the top-level duration
            duration_span = header_area.find('span', title=re.compile(r'[Cc]ourse\s+duration', re.I))
            if duration_span:
                duration_text = duration_span.get_text(strip=True)
                duration_match = re.search(r'(\d+\s*(?:hour|hr|h)(?:\s+\d+\s*(?:min|m))?|\d+\s*(?:min|m))', duration_text, re.IGNORECASE)
                if duration_match:
                    top_duration = duration_match.group(1).strip()
            
            # Look for "X Lessons" text in the same area (sibling spans)
            # Find all spans in the header area and look for lessons count
            all_spans = header_area.find_all('span')
            for span in all_spans:
                span_text = span.get_text(strip=True)
                lesson_match = re.search(r'(\d+)\s*Lessons?', span_text, re.IGNORECASE)
                if lesson_match:
                    top_lesson_count = lesson_match.group(1)
                    break
            
            # Construct top-level text
            if top_duration or top_lesson_count:
                parts = []
                if top_lesson_count:
                    parts.append(f"{top_lesson_count} Lessons")
                if top_duration:
                    parts.append(top_duration)
                top_level_text = ' '.join(parts)
    
    # If we didn't find top-level in header, try the first duration element in main content
    if not top_duration:
        all_duration_elems = soup.find_all('p', class_=lambda x: x and 'text-muted' in x and 'mt-1' in x)
        if all_duration_elems:
            # The first one might be top-level (but could also be first section)
            top_level_duration_elem = all_duration_elems[0]
            top_level_text = top_level_duration_elem.get_text(strip=True)
            
            # Extract time from top-level
            top_duration_match = re.search(r'(\d+\s*(?:hour|hr|h)(?:\s+\d+\s*(?:min|m))?|\d+\s*(?:min|m))', top_level_text, re.IGNORECASE)
            top_duration = top_duration_match.group(1).strip() if top_duration_match else top_level_text
            
            # Extract lesson count from top-level
            top_lesson_match = re.search(r'(\d+)\s*Lesson', top_level_text, re.IGNORECASE)
            top_lesson_count = top_lesson_match.group(1) if top_lesson_match else None
    
    if verbose:
        print(f"✓ Top-level duration: {top_duration}, Lessons: {top_lesson_count}")
    
    # Extract section durations by finding h2 headings and their associated durations
    # Also extract lessons within each section
    sections = []
    h2_headings = soup.find_all('h2')
    
    for heading in h2_headings:
        heading_text = heading.get_text(strip=True)
        # Skip if it's not a section heading (like "Related Courses")
        if not heading_text or 'Related' in heading_text or len(heading_text) > 100:
            continue
        
        # Find duration in the parent container of this heading
        parent = heading.parent
        if parent:
            duration_elem = parent.find('p', class_=lambda x: x and 'text-muted' in x and 'mt-1' in x)
            if duration_elem:
                section_text = duration_elem.get_text(strip=True)
                # Check if it contains duration info
                if any(word in section_text.lower() for word in ['lesson', 'min', 'hour', 'h', 'm']):
                    # Extract duration
                    section_duration_match = re.search(r'(\d+\s*(?:hour|hr|h)(?:\s+\d+\s*(?:min|m))?|\d+\s*(?:min|m))', section_text, re.IGNORECASE)
                    section_duration = section_duration_match.group(1).strip() if section_duration_match else section_text
                    
                    # Extract lesson count
                    section_lesson_match = re.search(r'(\d+)\s*Lesson', section_text, re.IGNORECASE)
                    section_lesson_count = section_lesson_match.group(1) if section_lesson_match else None
                    
                    # Extract lessons from this section
                    lessons = []
                    # Look for ordered list (ol) with class "list-group" in the parent container
                    lesson_list = parent.find('ol', class_=lambda x: x and 'list-group' in str(x))
                    if lesson_list:
                        # Find all anchor tags that contain list items (lessons are wrapped in <a> tags)
                        lesson_links = lesson_list.find_all('a', href=True)
                        for lesson_link_elem in lesson_links:
                            lesson_item = lesson_link_elem.find('li', class_=lambda x: x and 'list-group-item' in str(x))
                            if not lesson_item:
                                continue
                            
                            # Extract lesson title - it's in a span with class "mx-1"
                            lesson_title = None
                            title_span = lesson_item.find('span', class_=lambda x: x and 'mx-1' in str(x) if x else False)
                            if title_span:
                                # Get all text from the span, but remove the number prefix
                                lesson_title = title_span.get_text(strip=True)
                                # Remove number prefix like "1. " or "2. "
                                lesson_title = re.sub(r'^\d+\.\s*', '', lesson_title)
                            else:
                                # Fallback: get all text and clean it
                                lesson_title = lesson_item.get_text(strip=True)
                                # Remove number prefix
                                lesson_title = re.sub(r'^\d+\.\s*', '', lesson_title)
                            
                            # Extract lesson duration (format: MM:SS or HH:MM:SS)
                            # Duration is in <small class="ml-auto text-muted">
                            lesson_duration = None
                            duration_small = lesson_item.find('small', class_=lambda x: x and 'ml-auto' in str(x) and 'text-muted' in str(x) if x else False)
                            if duration_small:
                                duration_text = duration_small.get_text(strip=True)
                                # Duration is in MM:SS or HH:MM:SS format
                                if re.match(r'\d{1,2}:\d{2}(:\d{2})?', duration_text):
                                    lesson_duration = duration_text
                            
                            # Extract lesson link
                            lesson_link = lesson_link_elem.get('href')
                            if lesson_link and lesson_link.startswith('/'):
                                lesson_link = f"https://realpython.com{lesson_link}"
                            
                            if lesson_title:
                                lessons.append({
                                    'lesson_title': lesson_title,
                                    'duration': lesson_duration,  # MM:SS or HH:MM:SS format
                                    'duration_minutes': parse_time_to_minutes(lesson_duration) if lesson_duration else 0,
                                    'link': lesson_link
                                })
                    
                    sections.append({
                        'section_name': heading_text,
                        'duration': section_duration,
                        'lesson_count': section_lesson_count,
                        'full_text': section_text,
                        'lessons': lessons
                    })
    
    # Remove duplicates (same section might appear multiple times in nested structures)
    seen_sections = {}
    unique_sections = []
    for section in sections:
        key = section['section_name']
        if key not in seen_sections:
            seen_sections[key] = True
            unique_sections.append(section)
    
    if verbose and unique_sections:
        print(f"✓ Found {len(unique_sections)} sections with durations:")
        for sec in unique_sections:
            lessons_count = len(sec.get('lessons', []))
            print(f"   - {sec['section_name']}: {sec['duration']} ({sec['lesson_count']} lessons)")
            if lessons_count > 0:
                print(f"     └─ {lessons_count} lessons with individual durations")
                for i, lesson in enumerate(sec['lessons'][:3], 1):  # Show first 3
                    print(f"        {i}. {lesson['lesson_title']}: {lesson['duration']}")
                if lessons_count > 3:
                    print(f"        ... and {lessons_count - 3} more lessons")
    
    result = {
        'duration': top_duration,  # Main duration (for backward compatibility)
        'lesson_count': top_lesson_count,
        'full_text': top_level_text,
        'top_level_duration': top_duration,
        'top_level_lesson_count': top_lesson_count,
        'sections': unique_sections
    }
    
    return result

def parse_time_to_minutes(time_str):
    """Convert time string in MM:SS or HH:MM:SS format to total minutes"""
    if not time_str:
        return 0
    
    # Handle MM:SS or HH:MM:SS format
    time_match = re.match(r'(\d{1,2}):(\d{2})(?::(\d{2}))?', time_str)
    if time_match:
        if time_match.group(3):  # HH:MM:SS format
            hours = int(time_match.group(1))
            minutes = int(time_match.group(2))
            seconds = int(time_match.group(3))
            return hours * 60 + minutes + (seconds / 60)
        else:  # MM:SS format
            minutes = int(time_match.group(1))
            seconds = int(time_match.group(2))
            return minutes + (seconds / 60)
    
    return 0

def parse_duration_to_minutes(duration_str):
    """Convert duration string (e.g., '29m', '1h 30m', '2h') to total minutes"""
    if not duration_str:
        return 0
    
    total_minutes = 0
    # Pattern to match hours and minutes
    # Handle formats like "1h 30m", "2h", "29m", "1 hour 30 min"
    hour_match = re.search(r'(\d+)\s*(?:hour|hr|h)(?:\s|$)', duration_str, re.IGNORECASE)
    minute_match = re.search(r'(\d+)\s*(?:min|m)(?!\w)', duration_str, re.IGNORECASE)
    
    if hour_match:
        total_minutes += int(hour_match.group(1)) * 60
    if minute_match:
        total_minutes += int(minute_match.group(1))
    
    return total_minutes

def extract_all_resources_from_path(soup, path_title):
    """Extract all resources from a learning path page with proper categorization"""
    if not soup:
        return []
    
    all_containers = soup.find_all('div', class_=lambda x: x and 'container' in x and 'border' in x and 'rounded' in x and 'shadow-sm' in x)
    
    resources = []
    for container in all_containers:
        # Check if this container has a link to a course or quiz
        link_elem = container.find('a', class_='stretched-link', href=True)
        if not link_elem:
            continue
        
        href = link_elem.get('href', '')
        if not ('/courses/' in href or '/quizzes/' in href or '/tutorials/' in href):
            continue
        
        # Extract resource data
        resource_type = None
        type_elem = container.find('p', class_=lambda x: x and 'small' in x and 'text-muted' in x)
        if type_elem:
            strong_elem = type_elem.find('strong')
            if strong_elem:
                resource_type = strong_elem.get_text(strip=True)
        
        title = None
        title_elem = container.find('h3', class_=lambda x: x and 'h4' in x)
        if title_elem:
            link_in_h3 = title_elem.find('a')
            if link_in_h3:
                title = link_in_h3.get_text(strip=True)
            else:
                title = title_elem.get_text(strip=True)
        
        link = href
        if link and link.startswith('/'):
            link = f"https://realpython.com{link}"
        
        description = None
        desc_elem = container.find('p', class_=lambda x: x and 'small' in x and 'mb-0' in x and not ('text-muted' in x if x else False))
        if desc_elem:
            desc_text = desc_elem.get_text(strip=True)
            if desc_text:
                description = desc_text
        
        # Determine category - check title FIRST (most reliable indicator)
        # Exercises often have "Exercises:" in title but are under /courses/ URLs
        title_lower = (title or '').lower()
        category = 'other'
        
        if title_lower.startswith('exercises:') or 'exercise:' in title_lower or (title_lower.startswith('exercise') and ':' in title_lower):
            category = 'exercise'
        elif '/quizzes/' in href or (resource_type and 'quiz' in resource_type.lower()):
            category = 'quiz'
        elif '/courses/' in href or (resource_type and 'course' in resource_type.lower()):
            category = 'course'
        elif 'exercise' in href.lower() or (resource_type and 'exercise' in resource_type.lower()):
            category = 'exercise'
        
        resources.append({
            'category': category,
            'type': resource_type,
            'title': title,
            'description': description,
            'link': link,
            'duration': None,  # Will be filled when we visit the page
            'lesson_count': None,
            'duration_minutes': 0,
            'has_duration': category in ['course', 'exercise']  # Quizzes typically don't have duration
        })
    
    return resources

def process_single_learning_path(path, analyze_structure=False):
    """Process a single learning path in detail"""
    print(f"\n{'='*60}")
    print(f"PROCESSING: {path['title']}")
    print(f"{'='*60}")
    
    # Fetch learning path page
    path_soup = fetch_page(path['url'])
    if not path_soup:
        print(f"⚠️ Could not fetch learning path: {path['title']}")
        return []
    
    # Extract all resources
    resources = extract_all_resources_from_path(path_soup, path['title'])
    
    # Categorize for display
    courses = [r for r in resources if r['category'] == 'course']
    quizzes = [r for r in resources if r['category'] == 'quiz']
    exercises = [r for r in resources if r['category'] == 'exercise']
    other = [r for r in resources if r['category'] == 'other']
    
    print(f"  - {len(courses)} courses")
    print(f"  - {len(quizzes)} quizzes")
    print(f"  - {len(exercises)} exercises")
    if other:
        print(f"  - {len(other)} other resources")
    
    # Visit each resource page to get duration (only for courses and exercises)
    for res_idx, resource in enumerate(resources, 1):
        if not resource.get('link'):
            continue
        
        # Only fetch duration for courses and exercises (quizzes typically don't have duration)
        if resource['has_duration']:
            print(f"  [{res_idx}/{len(resources)}] {resource['category'].upper()}: {resource['title']}")
            duration_data = extract_duration_from_resource_page(resource['link'], verbose=False)
            
            if duration_data:
                resource['duration'] = duration_data.get('duration')
                resource['lesson_count'] = duration_data.get('lesson_count')
                resource['duration_minutes'] = parse_duration_to_minutes(resource['duration'])
                
                # Store section information if available
                resource['top_level_duration'] = duration_data.get('top_level_duration')
                resource['top_level_lesson_count'] = duration_data.get('top_level_lesson_count')
                resource['sections'] = duration_data.get('sections', [])
                
                # Calculate totals and validate
                if resource['sections']:
                    # Calculate section totals from lessons
                    for section in resource['sections']:
                        lessons = section.get('lessons', [])
                        if lessons:
                            lesson_total_minutes = sum(lesson.get('duration_minutes', 0) for lesson in lessons)
                            section['lessons_total_minutes'] = lesson_total_minutes
                            section_duration_minutes = parse_duration_to_minutes(section.get('duration', ''))
                            
                    # Calculate total from sections
                    section_total_minutes = sum(parse_duration_to_minutes(sec.get('duration', '')) for sec in resource['sections'])
                    resource['sections_total_minutes'] = section_total_minutes
                    print(f"      ✓ {resource['duration']} ({len(resource['sections'])} sections)")
                else:
                    print(f"      ✓ {resource['duration']}")
            else:
                print(f"      ⚠️ Could not extract duration")
        else:
            pass  # Skip quizzes (no duration expected)
        
        # No delay needed - working with local cache
    
    # Add path info to each resource
    for resource in resources:
        resource['learning_path'] = path['title']
        resource['learning_path_url'] = path['url']
    
    # Summary for this path
    total_minutes = sum(r.get('duration_minutes', 0) for r in resources)
    total_hours = total_minutes / 60
    
    print(f"✓ Complete: {len(resources)} resources, {total_hours:.1f}h total")
    
    return resources

def process_all_learning_paths(learning_paths, max_paths=None):
    """Process all learning paths one by one"""
    all_data = []
    
    paths_to_process = learning_paths[:max_paths] if max_paths else learning_paths
    
    for idx, path in enumerate(paths_to_process, 1):
        print(f"\n[{idx}/{len(paths_to_process)}] ", end="")
        resources = process_single_learning_path(path, analyze_structure=False)
        all_data.extend(resources)
    
    return all_data

def write_to_excel(data, learning_paths_order=None, filename='learning_paths_analysis.xlsx'):
    """Write all data to an Excel file in a table of contents format"""
    # Ensure output directory exists
    OUTPUT_DIR.mkdir(exist_ok=True)
    filepath = OUTPUT_DIR / filename
    print(f"\n{'='*60}")
    print(f"Writing data to Excel: {filename}")
    print(f"{'='*60}")
    
    wb = Workbook()
    
    # Remove default sheet, we'll create our own
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])
    
    # Define borders
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Style for headers
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    
    # Fonts for different levels
    path_font = Font(bold=True, size=13)
    course_font = Font(bold=True, size=12)
    section_font = Font(bold=True, size=11)
    lesson_font = Font(size=11)
    
    # Background colors for hierarchy
    path_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    course_fill = PatternFill(start_color="E7F3FF", end_color="E7F3FF", fill_type="solid")
    section_fill = PatternFill(start_color="F0F8FF", end_color="F0F8FF", fill_type="solid")
    lesson_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    # First, calculate all data we need before creating sheets
    # Group data by learning path
    paths_data = defaultdict(list)
    for resource in data:
        path = resource.get('learning_path', 'Unknown')
        paths_data[path].append(resource)
    
    # Determine order: use provided order if available, otherwise use sorted order
    if learning_paths_order:
        # Create ordered list preserving HTML order
        ordered_paths = []
        seen_paths = set()
        for path_info in learning_paths_order:
            path_title = path_info.get('title') if isinstance(path_info, dict) else path_info
            if path_title in paths_data and path_title not in seen_paths:
                ordered_paths.append(path_title)
                seen_paths.add(path_title)
        # Add any paths not in the original order (shouldn't happen, but just in case)
        for path_title in paths_data.keys():
            if path_title not in seen_paths:
                ordered_paths.append(path_title)
    else:
        # Fallback to sorted order if no order provided
        ordered_paths = sorted(paths_data.keys())
    
    # Calculate path_totals
    path_totals = {}
    for learning_path in ordered_paths:
        resources = paths_data[learning_path]
        total_resources = len(resources)
        total_minutes = sum(r.get('duration_minutes', 0) for r in resources)
        path_totals[learning_path] = {
            'resource_count': total_resources,
            'total_minutes': total_minutes
        }
    
    # Create Learning Paths table of contents sheet FIRST (index 0)
    ws_main = wb.create_sheet("Learning Paths", 0)
    ws_main.title = "Learning Paths"
    
    # Create main sheet with list of all learning paths
    main_headers = ['#', 'Learning Path', 'Resources', 'Total Duration', 'Total Duration (Minutes)', 
                    'Completed Resources', 'Completed Minutes', 'Progress %', 'Progress Bar']
    for col_idx, header in enumerate(main_headers, 1):
        cell = ws_main.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    
    # First, create all sheets (empty) so hyperlinks can reference them
    sheet_name_map = {}  # Map learning_path -> sheet_name
    used_sheet_names = set()  # Track used sheet names to avoid duplicates
    for learning_path in ordered_paths:
        # Create safe sheet name (Excel sheet names have limitations)
        base_sheet_name = learning_path[:31]  # Excel sheet name max 31 chars
        base_sheet_name = base_sheet_name.replace('/', '-').replace('\\', '-').replace('?', '').replace('*', '').replace('[', '').replace(']', '').replace(':', '')
        
        # Handle duplicate sheet names by appending a number
        sheet_name = base_sheet_name
        counter = 1
        while sheet_name in used_sheet_names:
            # Append counter, but keep within 31 char limit
            suffix = f"_{counter}"
            max_base_len = 31 - len(suffix)
            sheet_name = base_sheet_name[:max_base_len] + suffix
            counter += 1
        
        used_sheet_names.add(sheet_name)
        sheet_name_map[learning_path] = sheet_name
        # Create empty sheet first
        ws = wb.create_sheet(sheet_name)
        ws.title = sheet_name
    
    # Write learning paths to main sheet with hyperlinks (now that sheets exist)
    main_row = 2
    for path_idx, learning_path in enumerate(ordered_paths, 1):
        resources = paths_data[learning_path]
        # Get totals from path_totals (already calculated)
        totals = path_totals[learning_path]
        total_resources = totals['resource_count']
        total_minutes = totals['total_minutes']
        total_duration = f"{int(total_minutes // 60)}h {int(total_minutes % 60)}m" if total_minutes >= 60 else f"{int(total_minutes)}m"
        
        # Get sheet name from map
        sheet_name = sheet_name_map[learning_path]
        
        # Write to main sheet
        ws_main.cell(row=main_row, column=1, value=path_idx).border = thin_border
        cell = ws_main.cell(row=main_row, column=2, value=learning_path)
        cell.font = path_font
        cell.border = thin_border
        # Add internal hyperlink (sheet now exists)
        # Format: "# 'SheetName'!CellReference" for internal references
        # Use quotes around sheet name if it contains spaces or special chars
        cell.hyperlink = f"# '{sheet_name}'!A1"
        cell.font = Font(bold=True, size=13, color="0563C1", underline="single")
        ws_main.cell(row=main_row, column=3, value=total_resources).border = thin_border
        ws_main.cell(row=main_row, column=4, value=total_duration).border = thin_border
        ws_main.cell(row=main_row, column=5, value=total_minutes).border = thin_border
        
        # Column F: Completed Resources (FORMULA - counts only "Yes" in LESSON rows)
        # Count only rows where Row Type (column E) = "Lesson" and Completed (column D) = "Yes"
        completed_formula = f"=COUNTIFS('{sheet_name}'!E:E,\"Lesson\",'{sheet_name}'!D:D,\"Yes\")"
        cell = ws_main.cell(row=main_row, column=6, value=completed_formula)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(size=11)
        
        # Column G: Completed Minutes (FORMULA - sums minutes where Row Type="Lesson" and Completed="Yes")
        # Use SUMIFS to count only lesson rows
        completed_minutes_formula = f"=SUMIFS('{sheet_name}'!C:C,'{sheet_name}'!E:E,\"Lesson\",'{sheet_name}'!D:D,\"Yes\")"
        cell = ws_main.cell(row=main_row, column=7, value=completed_minutes_formula)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(size=11)
        
        # Column H: Progress % (FORMULA - calculates percentage based on minutes)
        progress_formula = f"=IF(E{main_row}>0,G{main_row}/E{main_row},0)"
        cell = ws_main.cell(row=main_row, column=8, value=progress_formula)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(size=11)
        cell.number_format = '0.0%'  # Standard percentage format
        
        # Column I: Progress Bar (same as progress %, used for visualization)
        # For 0%, return empty string - empty cells don't show data bars
        progress_bar_formula = f"=IF(H{main_row}=0,\"\",H{main_row}*100)"
        cell = ws_main.cell(row=main_row, column=9, value=progress_bar_formula)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(size=11)
        cell.number_format = '0.0"%"'
        
        main_row += 1
    
    # Add grand total row
    grand_total_row = main_row + 1
    total_resources_all = sum(t['resource_count'] for t in path_totals.values())
    total_minutes_all = sum(t['total_minutes'] for t in path_totals.values())
    total_duration_all = f"{int(total_minutes_all // 60)}h {int(total_minutes_all % 60)}m" if total_minutes_all >= 60 else f"{int(total_minutes_all)}m"
    
    # Grand total row
    cell = ws_main.cell(row=grand_total_row, column=1, value="GRAND TOTAL")
    cell.font = Font(bold=True, size=13)
    cell.border = thin_border
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    
    ws_main.cell(row=grand_total_row, column=2, value="").border = thin_border  # Empty for Learning Path column
    cell = ws_main.cell(row=grand_total_row, column=3, value=total_resources_all)
    cell.font = Font(bold=True)
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center", vertical="center")
    
    cell = ws_main.cell(row=grand_total_row, column=4, value=total_duration_all)
    cell.font = Font(bold=True)
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center", vertical="center")
    
    cell = ws_main.cell(row=grand_total_row, column=5, value=total_minutes_all)
    cell.font = Font(bold=True)
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Grand Total - Completed Resources
    cell = ws_main.cell(row=grand_total_row, column=6, value=f"=SUM(F2:F{main_row-1})")
    cell.font = Font(bold=True)
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Grand Total - Completed Minutes
    cell = ws_main.cell(row=grand_total_row, column=7, value=f"=SUM(G2:G{main_row-1})")
    cell.font = Font(bold=True)
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Grand Total - Progress %
    cell = ws_main.cell(row=grand_total_row, column=8, value=f"=IF(E{grand_total_row}>0,G{grand_total_row}/E{grand_total_row},0)")
    cell.font = Font(bold=True)
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.number_format = '0.0%'  # Standard percentage format
    
    # Grand Total - Progress Bar
    cell = ws_main.cell(row=grand_total_row, column=9, value=f"=IF(H{grand_total_row}=0,\"\",H{grand_total_row}*100)")
    cell.font = Font(bold=True)
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.number_format = '0.0"%"'
    
    # Add Progress Bars (Data Bars) to column I (Progress Bar column)
    progress_bar_rule = DataBarRule(
        start_type="num", start_value=0,
        end_type="num", end_value=100,
        color="63BE7B",  # Green
        showValue=True, minLength=0, maxLength=None
    )
    ws_main.conditional_formatting.add(f"I2:I{grand_total_row}", progress_bar_rule)
    
    # Add Pie Chart (based on total hours/minutes, not resource count)
    chart_row = grand_total_row + 3
    
    # Chart data labels
    ws_main.cell(row=chart_row, column=11, value="Chart Data").font = Font(bold=True, size=11)
    ws_main.cell(row=chart_row, column=12, value="Hours").font = Font(bold=True, size=11)
    
    # Labels
    chart_data_start = chart_row + 1
    ws_main.cell(row=chart_data_start, column=11, value="Completed")
    ws_main.cell(row=chart_data_start + 1, column=11, value="Remaining")
    
    # Values (formulas)
    # Completed hours: completed minutes / 60
    ws_main.cell(row=chart_data_start, column=12, value=f"=ROUND(G{grand_total_row}/60,1)")
    # Remaining hours: (total minutes - completed minutes) / 60
    ws_main.cell(row=chart_data_start + 1, column=12, value=f"=ROUND((E{grand_total_row}-G{grand_total_row})/60,1)")
    
    # Create pie chart
    pie = PieChart()
    pie.title = "Overall Learning Progress (Hours)"
    pie.height = 12
    pie.width = 15
    
    # Data reference: values in column L (12), rows chart_data_start to chart_data_start+1
    data = Reference(ws_main, min_col=12, min_row=chart_data_start, max_col=12, max_row=chart_data_start + 1)
    # Labels reference: labels in column K (11), rows chart_data_start to chart_data_start+1
    labels = Reference(ws_main, min_col=11, min_row=chart_data_start, max_col=11, max_row=chart_data_start + 1)
    
    pie.add_data(data, titles_from_data=False)
    pie.set_categories(labels)
    
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showValue = True
    
    # Position chart more to the right (column L, row 2)
    ws_main.add_chart(pie, "L2")
    
    # Auto-adjust main sheet column widths
    ws_main.column_dimensions['A'].width = 8
    ws_main.column_dimensions['B'].width = 60
    ws_main.column_dimensions['C'].width = 12
    ws_main.column_dimensions['D'].width = 18
    ws_main.column_dimensions['E'].width = 22
    ws_main.column_dimensions['F'].width = 18  # Completed Resources
    ws_main.column_dimensions['G'].width = 18  # Completed Minutes
    ws_main.column_dimensions['H'].width = 12  # Progress %
    ws_main.column_dimensions['I'].width = 25  # Progress Bar (expanded for better visibility)
    
    # Populate individual sheets for each learning path
    for learning_path in ordered_paths:
        resources = paths_data[learning_path]
        # Get sheet name from map (sheet already exists)
        sheet_name = sheet_name_map[learning_path]
        ws = wb[sheet_name]  # Get existing sheet instead of creating new one
        
        # Headers - add "Completed" column and "Row Type" helper column
        headers = ['Title', 'Duration', 'Duration (Minutes)', 'Completed', 'Row Type']
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        
        # Add data validation for Completed column
        completed_dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
        completed_dv.error = "Please enter Yes or No"
        completed_dv.errorTitle = "Invalid Entry"
        ws.add_data_validation(completed_dv)
        
        row = 2
        course_num = 0
        last_data_row = row  # Track last row with data for pie chart
        all_lesson_rows = []  # Track all lesson rows for main sheet counting
        
        for resource_idx, resource in enumerate(resources):
            course_num += 1
            title = resource.get('title', '')
            category = resource.get('category', '').upper()
            duration = resource.get('top_level_duration') or resource.get('duration', '')
            duration_minutes = resource.get('duration_minutes', 0)
            
            # Store course row for later formula application
            course_row = row
            
            # Course row with numbering
            course_title = f"    {course_num}. {title} ({category})"
            for col in range(1, 4):
                cell = ws.cell(row=row, column=col)
                cell.font = course_font
                cell.fill = course_fill
                cell.border = thin_border
                if col == 1:
                    cell.value = course_title
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif col == 2:
                    cell.value = duration
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col == 3:
                    cell.value = duration_minutes
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Column D: Completed - will be set later (formula or manual)
            # Don't set value yet - will be set after we know if it has sections
            completed_cell = ws.cell(row=row, column=4)
            completed_cell.border = thin_border
            completed_cell.alignment = Alignment(horizontal="center", vertical="center")
            completed_cell.font = course_font
            completed_cell.fill = course_fill
            
            # Column E: Row Type (helper for main sheet counting)
            ws.cell(row=row, column=5, value="Course").border = thin_border
            ws.cell(row=row, column=5).font = course_font
            ws.cell(row=row, column=5).fill = course_fill
            ws.cell(row=row, column=5).alignment = Alignment(horizontal="center", vertical="center")
            
            last_data_row = row
            row += 1
            
            # Track sections and their lesson rows for this course
            section_rows = []
            section_lesson_rows = {}  # Dict: section_index -> [lesson_rows]
            
            # Write sections with numbering
            sections = resource.get('sections', [])
            if sections:
                section_num = 0
                for section in sections:
                    section_num += 1
                    section_name = section.get('section_name', '')
                    section_duration = section.get('duration', '')
                    section_duration_minutes = parse_duration_to_minutes(section_duration)
                    
                    # Store section row for later formula application
                    section_row = row
                    section_rows.append(section_row)
                    
                    # Section with numbering (course.section format)
                    section_title = f"        {course_num}.{section_num} {section_name}"
                    for col in range(1, 4):
                        cell = ws.cell(row=row, column=col)
                        cell.font = section_font
                        cell.fill = section_fill
                        cell.border = thin_border
                        if col == 1:
                            cell.value = section_title
                            cell.alignment = Alignment(horizontal="left", vertical="center")
                        elif col == 2:
                            cell.value = section_duration
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                        elif col == 3:
                            cell.value = section_duration_minutes
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                    
                    # Column D: Completed - will be set later (formula or manual)
                    # Don't set value yet - will be set after we know if it has lessons
                    completed_cell = ws.cell(row=row, column=4)
                    completed_cell.border = thin_border
                    completed_cell.alignment = Alignment(horizontal="center", vertical="center")
                    completed_cell.font = section_font
                    completed_cell.fill = section_fill
                    
                    # Column E: Row Type
                    ws.cell(row=row, column=5, value="Section").border = thin_border
                    ws.cell(row=row, column=5).font = section_font
                    ws.cell(row=row, column=5).fill = section_fill
                    ws.cell(row=row, column=5).alignment = Alignment(horizontal="center", vertical="center")
                    
                    last_data_row = row
                    row += 1
                    
                    # Track lesson rows for this section
                    lesson_rows = []
                    
                    # Write lessons with numbering
                    lessons = section.get('lessons', [])
                    if lessons:
                        lesson_num = 0
                        for lesson in lessons:
                            lesson_num += 1
                            lesson_title = lesson.get('lesson_title', '')
                            lesson_duration = lesson.get('duration', '')
                            lesson_duration_minutes = lesson.get('duration_minutes', 0)
                            
                            # Lesson with numbering (course.section.lesson format)
                            lesson_title_formatted = f"            {course_num}.{section_num}.{lesson_num} {lesson_title}"
                            for col in range(1, 4):
                                cell = ws.cell(row=row, column=col)
                                cell.font = lesson_font
                                cell.fill = lesson_fill
                                cell.border = thin_border
                                if col == 1:
                                    cell.value = lesson_title_formatted
                                    cell.alignment = Alignment(horizontal="left", vertical="center")
                                elif col == 2:
                                    cell.value = lesson_duration
                                    cell.alignment = Alignment(horizontal="center", vertical="center")
                                elif col == 3:
                                    cell.value = round(lesson_duration_minutes, 2)
                                    cell.alignment = Alignment(horizontal="center", vertical="center")
                            
                            # Column D: Completed - MANUAL DROPDOWN (only for lessons)
                            completed_cell = ws.cell(row=row, column=4, value="No")
                            completed_cell.border = thin_border
                            completed_cell.alignment = Alignment(horizontal="center", vertical="center")
                            completed_cell.font = lesson_font
                            completed_cell.fill = lesson_fill
                            completed_dv.add(f"D{row}")  # Manual dropdown for lessons
                            
                            # Column E: Row Type
                            ws.cell(row=row, column=5, value="Lesson").border = thin_border
                            ws.cell(row=row, column=5).font = lesson_font
                            ws.cell(row=row, column=5).fill = lesson_fill
                            ws.cell(row=row, column=5).alignment = Alignment(horizontal="center", vertical="center")
                            
                            lesson_rows.append(row)
                            all_lesson_rows.append(row)  # Track for main sheet
                            last_data_row = row
                            row += 1
                    
                    # Store lesson rows for this section
                    section_lesson_rows[section_num] = lesson_rows
                    
                    # Apply formula to section if it has lessons, otherwise manual dropdown
                    if lesson_rows:
                        # Section has lessons - use formula
                        # Formula: If all lessons are "Yes", section is "Yes", otherwise "No"
                        lesson_range = f"D{lesson_rows[0]}:D{lesson_rows[-1]}"
                        # Count "Yes" values and compare to total number of lessons
                        # Use COUNTA to count all non-empty cells (should equal number of lessons)
                        formula = f'=IF(COUNTIF({lesson_range},"Yes")=COUNTA({lesson_range}),"Yes","No")'
                        cell = ws.cell(row=section_row, column=4, value=formula)
                        cell.border = thin_border
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.font = section_font
                        cell.fill = section_fill
                        # Formula cells don't need data validation
                    else:
                        # Section has no lessons - manual dropdown (starts as "No")
                        cell = ws.cell(row=section_row, column=4, value="No")
                        cell.border = thin_border
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.font = section_font
                        cell.fill = section_fill
                        completed_dv.add(f"D{section_row}")
            
            # Apply formula to course if it has sections, otherwise manual dropdown
            if section_rows:
                # Course has sections - use formula
                # Formula: If all sections are "Yes", course is "Yes", otherwise "No"
                section_range = f"D{section_rows[0]}:D{section_rows[-1]}"
                # Count "Yes" values and compare to total number of sections
                formula = f'=IF(COUNTIF({section_range},"Yes")=COUNTA({section_range}),"Yes","No")'
                cell = ws.cell(row=course_row, column=4, value=formula)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = course_font
                cell.fill = course_fill
                # Formula cells don't need data validation
            else:
                # Course has no sections - manual dropdown (starts as "No")
                cell = ws.cell(row=course_row, column=4, value="No")
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = course_font
                cell.fill = course_fill
                completed_dv.add(f"D{course_row}")
            
            # Add blank row after each course
            if resource_idx < len(resources) - 1:
                for col in range(1, 6):  # Include columns A-E
                    cell = ws.cell(row=row, column=col, value="")
                    cell.border = thin_border
                row += 1
        
        # Calculate total minutes for this learning path
        total_minutes = sum(r.get('duration_minutes', 0) for r in resources)
        
        # Add summary row with totals
        summary_row = last_data_row + 2
        ws.cell(row=summary_row, column=1, value="TOTAL").font = Font(bold=True, size=12)
        total_duration = f"{int(total_minutes // 60)}h {int(total_minutes % 60)}m" if total_minutes >= 60 else f"{int(total_minutes)}m"
        ws.cell(row=summary_row, column=2, value=total_duration).font = Font(bold=True)
        ws.cell(row=summary_row, column=3, value=total_minutes).font = Font(bold=True)
        ws.cell(row=summary_row, column=4, value="").font = Font(bold=True)
        ws.cell(row=summary_row, column=5, value="").font = Font(bold=True)  # Row Type column
        
        # Add pie chart for this learning path (positioned to the right)
        chart_row = summary_row + 3
        
        # Chart data labels
        ws.cell(row=chart_row, column=6, value="Chart Data").font = Font(bold=True, size=11)
        ws.cell(row=chart_row, column=7, value="Hours").font = Font(bold=True, size=11)
        
        # Labels
        chart_data_start = chart_row + 1
        ws.cell(row=chart_data_start, column=6, value="Completed")
        ws.cell(row=chart_data_start + 1, column=6, value="Remaining")
        
        # Values (formulas)
        # Completed hours: sum minutes where Row Type="Lesson" and Completed="Yes" / 60
        ws.cell(row=chart_data_start, column=7, value=f"=ROUND(SUMIFS(C2:C{last_data_row},E2:E{last_data_row},\"Lesson\",D2:D{last_data_row},\"Yes\")/60,1)")
        # Remaining hours: (total lesson minutes - completed lesson minutes) / 60
        # Total lesson minutes: sum of all lesson minutes
        ws.cell(row=chart_data_start + 1, column=7, value=f"=ROUND((SUMIF(E2:E{last_data_row},\"Lesson\",C2:C{last_data_row})-SUMIFS(C2:C{last_data_row},E2:E{last_data_row},\"Lesson\",D2:D{last_data_row},\"Yes\"))/60,1)")
        
        # Create pie chart
        pie = PieChart()
        pie.title = f"{learning_path} Progress (Hours)"
        pie.height = 10
        pie.width = 12
        
        # Data reference: values in column G (7), rows chart_data_start to chart_data_start+1
        data = Reference(ws, min_col=7, min_row=chart_data_start, max_col=7, max_row=chart_data_start + 1)
        # Labels reference: labels in column F (6), rows chart_data_start to chart_data_start+1
        labels = Reference(ws, min_col=6, min_row=chart_data_start, max_col=6, max_row=chart_data_start + 1)
        
        pie.add_data(data, titles_from_data=False)
        pie.set_categories(labels)
        
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.dataLabels.showValue = True
        
        # Position chart to the right (column I, row 2)
        ws.add_chart(pie, "I2")
        
        # Auto-adjust column widths for this sheet
        ws.column_dimensions['A'].width = 80
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 12  # Completed column
        ws.column_dimensions['E'].width = 12  # Row Type column (helper, can be hidden)
    
    wb.save(filepath)
    print(f"✓ Excel file saved: {filepath}")
    print(f"✓ Total resources: {len(data)}")
    print(f"✓ Total learning paths: {len(path_totals)}")
    total_minutes_all = sum(t['total_minutes'] for t in path_totals.values())
    print(f"✓ Total duration: {total_minutes_all} minutes ({total_minutes_all/60:.1f} hours)")
    print(f"\n📊 PROGRESS TRACKING FEATURES:")
    print(f"  ✓ Progress tracking columns in main sheet")
    print(f"  ✓ Progress bars (data bars) for visual progress")
    print(f"  ✓ Overall pie chart showing total progress")
    print(f"  ✓ Individual pie charts in each learning path sheet")
    print(f"  ✓ 'Completed' column - ONLY LESSONS have manual dropdowns")
    print(f"  ✓ Sections and courses auto-complete when all children are done")
    print(f"\n💡 HOW TO USE:")
    print(f"  1. Open the Excel file")
    print(f"  2. Go to any learning path sheet")
    print(f"  3. Mark LESSONS as 'Yes' when completed (only editable level)")
    print(f"  4. Sections and courses update automatically based on their lessons")
    print(f"  5. Progress updates automatically in main sheet and individual sheets!")

if __name__ == "__main__":
    # Step 1: Get all learning paths
    print("=" * 60)
    print("REAL PYTHON LEARNING PATHS SCRAPER")
    print("=" * 60)
    
    # Check if we have cached HTML files
    if HTML_DIR.exists():
        main_file = HTML_DIR / "main_learning_paths.html"
        if main_file.exists():
            print(f"✓ Using cached HTML files from: {HTML_DIR.absolute()}")
            print("  (Run 'python download_all_html.py' to refresh cache)")
        else:
            print("⚠️ HTML cache directory exists but main file not found")
            print("  Fetching from web...")
    else:
        print("⚠️ No HTML cache found. Fetching from web...")
        print(f"  (Run 'python download_all_html.py' once to cache all HTML files)")
    
    soup = fetch_main_page()
    if not soup:
        print("⚠️ Could not fetch main page")
        exit(1)
    
    learning_paths = extract_learning_paths(soup)

    print(f"\n✓ Found {len(learning_paths)} learning paths")
    
    if not learning_paths:
        print("\n⚠️ No learning paths found")
        exit(1)
    
    # Process all learning paths
    print(f"\n{'='*60}")
    print(f"PROCESSING ALL LEARNING PATHS")
    print(f"{'='*60}")
    
    all_resources = process_all_learning_paths(learning_paths)
    
    if all_resources:
        print(f"\n{'='*60}")
        print(f"✓ Successfully processed {len(learning_paths)} learning paths!")
        print(f"✓ Total resources: {len(all_resources)}")
        print(f"{'='*60}")
        
        # Write to Excel (pass learning_paths to preserve order)
        write_to_excel(all_resources, learning_paths_order=learning_paths, filename='learning_paths_analysis.xlsx')
    else:
        print("\n⚠️ No resources found")
